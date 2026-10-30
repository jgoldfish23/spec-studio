#!/usr/bin/env python3
"""Build the golden-deal spreadsheet and the expected values the engine must match.

The spreadsheet is the independent check on the engine's money math. Every schedule value is
an Excel formula implementing the Reference cash-flow algorithm (plan, "Reference: cash-flow
algorithm", steps 1-15), working in dollars with ROUND(..., 2) at exactly the points where the
Reference says `round`. This script only lays the formulas out; the one thing it decides is the
number of period rows (N = ceil(saleClose)), which the sheet re-derives and checks.

Steps:
1. Read src/engine/__fixtures__/golden-deal.inputs.json (money in cents).
2. Write docs/golden-deal.xlsx: an `Inputs` sheet (labeled cells, money in dollars) and one
   schedule sheet per scenario (`base`, `stress`).
3. Recalculate with headless LibreOffice (throwaway profile, temp output dir) and read the
   recalculated values back with openpyxl.
4. Write src/engine/__fixtures__/golden-deal.expected.json, converting dollars to integer cents
   with round-half-away-from-zero of (value x 100).

Run with `npm run golden`.
"""

from __future__ import annotations

import json
import math
import shutil
import subprocess
import tempfile
from itertools import accumulate
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "src" / "engine" / "__fixtures__"
INPUTS_JSON = FIXTURES / "golden-deal.inputs.json"
EXPECTED_JSON = FIXTURES / "golden-deal.expected.json"
XLSX = ROOT / "docs" / "golden-deal.xlsx"

MONTH_DAYS = 365 / 12
SOFT_CODES = ("plans", "permits", "surveySoils")

# Stress: 93% of the base price, a very long 400 days on market, and the downside 5% overrun.
STRESS = {"priceCents": 125_550_000, "domDays": 400, "overrunPct": 0.05}

MONEY = "#,##0.00"
MONTHS = "0.000000"
RATE = "0.0000"
BOLD = Font(bold=True)

# Schedule sheet layout: summary block in A:B, period table from TABLE_FIRST_ROW down.
TABLE_HEADER_ROW = 64
TABLE_FIRST_ROW = TABLE_HEADER_ROW + 1


def dollars(cents: int) -> float:
    return cents / 100


def to_cents(value: object, what: str) -> int:
    """Dollars from the sheet -> integer cents, half away from zero. Rejects non-cent amounts."""
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise SystemExit(f"{what}: expected a number from the sheet, got {value!r}")
    x = value * 100
    magnitude = math.floor(abs(x) + 0.5)
    if abs(abs(x) - magnitude) > 1e-6:
        raise SystemExit(f"{what}: {value!r} dollars is not a whole number of cents")
    return -magnitude if x < 0 else magnitude


def period_count(timeline: dict, dom_days: float) -> int:
    """N = ceil(saleClose): the only thing this script decides (how many period rows to write)."""
    completion = timeline["preConMonths"] + timeline["buildMonths"]
    list_at = completion - timeline["listEarlyMonths"]
    contract = list_at + dom_days / MONTH_DAYS
    sale_close = max(contract + timeline["closingDays"] / MONTH_DAYS, completion)
    return math.ceil(sale_close)


# ---------------------------------------------------------------------------------------------
# Inputs sheet
# ---------------------------------------------------------------------------------------------


def write_inputs_sheet(wb: Workbook, d: dict, scenarios: list[tuple[str, dict]]) -> dict[str, str]:
    """Writes the Inputs sheet with workbook-level names. Returns scenario name -> row number."""
    ws = wb.active
    assert ws is not None
    ws.title = "Inputs"
    ws["A1"] = "Golden deal - inputs"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Money in dollars, rates as decimals. Source: src/engine/__fixtures__/golden-deal.inputs.json"

    row = 4

    def section(title: str) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = BOLD
        row += 1

    def value(label: str, name: str | None, v: object, fmt: str | None = None, note: str = "") -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=v)
        if fmt:
            cell.number_format = fmt
        if note:
            ws.cell(row=row, column=3, value=note)
        if name:
            wb.defined_names[name] = DefinedName(name, attr_text=f"Inputs!$B${row}")
        row += 1

    lot, tl, loan, carry, selling = d["lot"], d["timeline"], d["loan"], d["carry"], d["selling"]
    fin = lot["financing"]
    land_loan = fin["kind"] == "loan"

    section("Lot")
    value("Land price", "LandPrice", dollars(lot["priceCents"]), MONEY)
    value("Land closing costs", "LandClosing", dollars(lot["closingCents"]), MONEY)
    value("Land financing", None, fin["kind"])
    value("Land loan LTV", "LandLtv", fin["ltv"] if land_loan else 0, RATE, "0 when bought for cash")
    value("Land loan rate (annual)", "LandRate", fin["rate"] if land_loan else 0, RATE)
    value("Land loan points", "LandPointsPct", fin["points"] if land_loan else 0, RATE)
    value("Property tax rate (annual)", "TaxRate", lot["propertyTaxRate"], RATE)
    value("HOA per month", "HoaMonthly", dollars(lot["hoaMonthlyCents"]), MONEY)
    row += 1

    section("Timeline (months from land close)")
    value("Pre-construction months (P)", "PreConMonths", tl["preConMonths"])
    value("Build months (B)", "BuildMonths", tl["buildMonths"])
    value("List early months (L)", "ListEarlyMonths", tl["listEarlyMonths"])
    value("Closing days after contract", "ClosingDays", tl["closingDays"])
    value("Days per month", "MonthDays", "=365/12", MONTHS)
    row += 1

    section("Construction loan")
    value("Loan-to-cost", "Ltc", loan["ltc"], RATE)
    value("Rate (annual)", "LoanRate", loan["rate"], RATE)
    value("Points", "PointsPct", loan["points"], RATE)
    value("Draw fee per draw", "DrawFee", dollars(loan["drawFeeCents"]), MONEY)
    value("Term months (from build start)", "TermMonths", loan["termMonths"])
    value("Extension fee", "ExtensionPct", loan["extensionPct"], RATE, "of the commitment, per extension")
    value("Extension months", "ExtensionMonths", loan["extensionMonths"])
    value("Interest mode", "InterestMode", loan["interestMode"], None, "paid | capitalized")
    row += 1

    section("Carry")
    value("Builder's risk insurance (annual, of estimate)", "InsuranceRate", carry["insuranceRate"], RATE)
    value("Utilities per month, during build", "UtilBuild", dollars(carry["utilitiesBuildMonthlyCents"]), MONEY)
    value("Utilities per month, after completion", "UtilAfter", dollars(carry["utilitiesAfterMonthlyCents"]), MONEY)
    value("Staging per month, while listed", "StagingMonthly", dollars(carry["stagingMonthlyCents"]), MONEY)
    value("Upkeep per month, after completion", "UpkeepMonthly", dollars(carry["upkeepMonthlyCents"]), MONEY)
    row += 1

    section("Selling")
    value("Listing commission", "ListingPct", selling["listingPct"], RATE)
    value("Buyer-agent commission", "BuyerPct", selling["buyerPct"], RATE)
    value("Closing and title", "ClosingPct", selling["closingPct"], RATE)
    value("Concessions", "Concessions", dollars(selling["concessionsCents"]), MONEY)
    row += 1

    section("Scenarios")
    for col, head in enumerate(("Scenario", "Sale price", "Days on market", "Overrun"), start=1):
        ws.cell(row=row, column=col, value=head).font = BOLD
    row += 1
    scenario_rows: dict[str, str] = {}
    for name, s in scenarios:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=dollars(s["priceCents"])).number_format = MONEY
        ws.cell(row=row, column=3, value=s["domDays"])
        ws.cell(row=row, column=4, value=s["overrunPct"]).number_format = RATE
        scenario_rows[name] = str(row)
        row += 1

    # Estimate: bids in E:G, summary underneath.
    ws["E4"] = "Estimate (bids)"
    ws["E4"].font = BOLD
    for col, head in zip("EFG", ("Code", "Division", "Amount")):
        ws[f"{col}5"] = head
        ws[f"{col}5"].font = BOLD
    first = 6
    for i, bid in enumerate(d["bids"]):
        r = first + i
        ws[f"E{r}"] = bid["code"]
        ws[f"F{r}"] = "soft" if bid["code"] in SOFT_CODES else "hard"
        ws[f"G{r}"] = dollars(bid["amountCents"])
        ws[f"G{r}"].number_format = MONEY
    last = first + len(d["bids"]) - 1
    r = last + 2

    def est(label: str, name: str, v: object, fmt: str = MONEY) -> None:
        nonlocal r
        ws[f"E{r}"] = label
        ws[f"G{r}"] = v
        ws[f"G{r}"].number_format = fmt
        wb.defined_names[name] = DefinedName(name, attr_text=f"Inputs!$G${r}")
        r += 1

    est("Bid subtotal", "BidSubtotal", f"=SUM(G{first}:G{last})")
    est("Contingency rate", "ContingencyPct", d["contingencyPct"], RATE)
    est("Contingency", "Contingency", "=ROUND(BidSubtotal*ContingencyPct,2)")
    est("Estimate total", "EstimateTotal", "=BidSubtotal+Contingency")
    est("Soft costs (division soft)", "SoftCost", f'=SUMIF(F{first}:F{last},"soft",G{first}:G{last})')
    est("Hard costs (incl. contingency)", "HardCost", "=EstimateTotal-SoftCost")

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 16
    return scenario_rows


# ---------------------------------------------------------------------------------------------
# Schedule sheets
# ---------------------------------------------------------------------------------------------

# Period table columns: (key, header, number format). Formulas are built in write_schedule_sheet.
COLUMNS: list[tuple[str, str, str | None]] = [
    ("k", "Period k", None),
    ("ovBuildClose", "Overlap build start->sale close", MONTHS),
    ("ovAll", "Overlap 0->sale close", MONTHS),
    ("ovBuild", "Overlap build start->completion", MONTHS),
    ("ovAfter", "Overlap completion->sale close", MONTHS),
    ("ovListed", "Overlap list->contract", MONTHS),
    ("sine", "Draw sine sin(pi(i+0.5)/B)", "0.000000000"),
    ("draw", "Hard-cost draw", MONEY),
    ("spend", "Spend x", MONEY),
    ("spendBefore", "Spend before (S)", MONEY),
    ("drawnBefore", "Loan drawn before", MONEY),
    ("eqFirst", "Equity first", MONEY),
    ("loanDraw", "Loan draw (ln)", MONEY),
    ("equitySpend", "Equity spend (eq)", MONEY),
    ("spendThrough", "Spend through period", MONEY),
    ("drawnThrough", "Loan drawn through period", MONEY),
    ("balanceStart", "Loan balance at start", MONEY),
    ("points", "Construction points", MONEY),
    ("drawFee", "Draw fee", MONEY),
    ("interest", "Construction interest", MONEY),
    ("interestPaid", "Interest paid", MONEY),
    ("interestCap", "Interest capitalized", MONEY),
    ("balanceEnd", "Loan balance at end", MONEY),
    ("landInterest", "Land loan interest", MONEY),
    ("extensionFee", "Extension fees", MONEY),
    ("tax", "Property tax", MONEY),
    ("insurance", "Insurance", MONEY),
    ("utilities", "Utilities", MONEY),
    ("staging", "Staging", MONEY),
    ("upkeep", "Upkeep", MONEY),
    ("hoa", "HOA", MONEY),
    ("landOut", "Land purchase (equity)", MONEY),
    ("equityOut", "Equity out", MONEY),
    ("cumEquity", "Cumulative equity out", MONEY),
    (None, "", None),  # spacer
    ("extJ", "Extension j", None),
    ("extPeriod", "Extension j fee period", None),
]
COL = {key: get_column_letter(i + 1) for i, (key, _, _) in enumerate(COLUMNS) if key}

# Summary block rows: (key or None for a section heading, label, number format).
SUMMARY: list[tuple[str | None, str, str | None]] = [
    (None, "Scenario inputs", None),
    ("SalePrice", "Sale price", MONEY),
    ("DomDays", "Days on market", None),
    ("Overrun", "Cost overrun", RATE),
    (None, "Times (months) - Reference step 1", None),
    ("BuildStart", "Build start = P", MONTHS),
    ("Completion", "Completion = P + B", MONTHS),
    ("ListMonth", "List = completion - L", MONTHS),
    ("Contract", "Contract = list + DOM / MONTH_DAYS", MONTHS),
    ("SaleClose", "Sale close = max(contract + closingDays / MONTH_DAYS, completion)", MONTHS),
    ("Maturity", "Loan maturity = build start + term", MONTHS),
    ("PeriodCount", "Periods N = ceil(sale close)", None),
    ("RowsInSchedule", "Period rows in this sheet", None),
    (None, "Costs and loan - Reference steps 2-8", None),
    ("SoftActual", "Soft actual = round(soft x (1 + overrun))", MONEY),
    ("HardActual", "Hard actual = round(hard x (1 + overrun))", MONEY),
    ("DrawWeightSum", "Sum of draw sines", "0.000000000"),
    ("Commitment", "Loan commitment = round(LTC x estimate total)", MONEY),
    ("EquityReq", "Equity requirement = estimate total - commitment", MONEY),
    ("ConstrPoints", "Construction points = round(points x commitment)", MONEY),
    ("LandLoan", "Land loan = round(LTV x land price)", MONEY),
    ("LandPoints", "Land points = round(points x land loan)", MONEY),
    ("ExtCount", "Extensions n", None),
    ("ExtFee", "Extension fee each = round(extension % x commitment)", MONEY),
    ("CapFlag", "Interest capitalized (1 = yes)", None),
    (None, "Sale (at sale close, booked on period N-1) - Reference step 11", None),
    ("Commissions", "Commissions = round(price x (listing + buyer))", MONEY),
    ("SaleClosing", "Closing = round(price x closing %)", MONEY),
    ("SaleConcessions", "Concessions", MONEY),
    ("NetProceeds", "Net = price - commissions - closing - concessions", MONEY),
    ("LoanDrawn", "Loan drawn", MONEY),
    ("CapInterest", "Capitalized interest", MONEY),
    ("Payoff", "Loan payoff = drawn + capitalized interest + land loan", MONEY),
    ("EquityReturn", "Equity return = net - payoff", MONEY),
    (None, "Totals and metrics - Reference steps 12-15", None),
    ("BuildSpent", "Build spent = soft actual + hard actual", MONEY),
    ("Financing", "Financing", MONEY),
    ("Carry", "Carry", MONEY),
    ("TotalProjectCost", "Total project cost", MONEY),
    ("SellingCosts", "Selling costs", MONEY),
    ("Profit", "Profit", MONEY),
    ("PeakCash", "Peak cash", MONEY),
    ("BreakEvenPrice", "Break-even sale price", MONEY),
    ("BalanceAtCompletion", "Loan balance after period completion - 1", MONEY),
    ("MonthlyBurn", "Monthly burn while listed", MONEY),
    ("MaxDomNoExt", "Max DOM without extension (blank = no loan)", None),
    (None, "Checks (each should be 0)", None),
    ("CheckRows", "N - period rows", None),
    ("CheckSpend", "Sum of spend - build spent", MONEY),
    ("CheckFunding", "Sum of equity spend + loan draws - build spent", MONEY),
    ("CheckFlows", "Equity return - sum of equity out - profit", MONEY),
]


def write_schedule_sheet(wb: Workbook, name: str, n: int, scenario_row: str) -> dict[str, str]:
    """Writes one scenario's schedule. Returns summary key -> cell address (e.g. 'B13')."""
    ws: Worksheet = wb.create_sheet(name)
    first, last = TABLE_FIRST_ROW, TABLE_FIRST_ROW + n - 1

    def local_name(key: str, ref: str) -> None:
        ws.defined_names[key] = DefinedName(key, attr_text=f"'{name}'!{ref}")

    def col(key: str) -> str:
        """Sheet-scoped name of a period column's range, e.g. spend -> colSpend."""
        return "col" + key[0].upper() + key[1:]

    for key, letter in COL.items():
        local_name(col(key), f"${letter}${first}:${letter}${last}")

    # Summary formulas.
    f: dict[str, str] = {
        "SalePrice": f"=Inputs!$B${scenario_row}",
        "DomDays": f"=Inputs!$C${scenario_row}",
        "Overrun": f"=Inputs!$D${scenario_row}",
        "BuildStart": "=PreConMonths",
        "Completion": "=BuildStart+BuildMonths",
        "ListMonth": "=Completion-ListEarlyMonths",
        "Contract": "=ListMonth+DomDays/MonthDays",
        "SaleClose": "=MAX(Contract+ClosingDays/MonthDays,Completion)",
        "Maturity": "=BuildStart+TermMonths",
        "PeriodCount": "=CEILING(SaleClose,1)",
        "RowsInSchedule": f"=COUNT({col('k')})",
        "SoftActual": "=ROUND(SoftCost*(1+Overrun),2)",
        "HardActual": "=ROUND(HardCost*(1+Overrun),2)",
        "DrawWeightSum": f"=SUM({col('sine')})",
        "Commitment": "=ROUND(Ltc*EstimateTotal,2)",
        "EquityReq": "=EstimateTotal-Commitment",
        "ConstrPoints": "=IF(Commitment>0,ROUND(PointsPct*Commitment,2),0)",
        "LandLoan": "=ROUND(LandLtv*LandPrice,2)",
        "LandPoints": "=ROUND(LandPointsPct*LandLoan,2)",
        "ExtCount": "=IF(AND(Commitment>0,SaleClose>Maturity),CEILING((SaleClose-Maturity)/ExtensionMonths,1),0)",
        "ExtFee": "=ROUND(ExtensionPct*Commitment,2)",
        "CapFlag": '=IF(InterestMode="capitalized",1,0)',
        "Commissions": "=ROUND(SalePrice*(ListingPct+BuyerPct),2)",
        "SaleClosing": "=ROUND(SalePrice*ClosingPct,2)",
        "SaleConcessions": "=Concessions",
        "NetProceeds": "=SalePrice-Commissions-SaleClosing-SaleConcessions",
        "LoanDrawn": f"=SUM({col('loanDraw')})",
        "CapInterest": f"=SUM({col('interestCap')})",
        "Payoff": "=LoanDrawn+CapInterest+LandLoan",
        "EquityReturn": "=NetProceeds-Payoff",
        "BuildSpent": "=SoftActual+HardActual",
        "Financing": (
            f"=SUM({col('points')})+LandPoints+SUM({col('interest')})+SUM({col('landInterest')})"
            f"+SUM({col('drawFee')})+SUM({col('extensionFee')})"
        ),
        "Carry": (
            f"=SUM({col('tax')})+SUM({col('insurance')})+SUM({col('utilities')})"
            f"+SUM({col('staging')})+SUM({col('upkeep')})+SUM({col('hoa')})"
        ),
        "TotalProjectCost": "=LandPrice+LandClosing+BuildSpent+Financing+Carry",
        "SellingCosts": "=Commissions+SaleClosing+SaleConcessions",
        "Profit": "=SalePrice-SellingCosts-TotalProjectCost",
        "PeakCash": f"=MAX({col('cumEquity')})",
        "BreakEvenPrice": "=ROUND((TotalProjectCost+Concessions)/(1-ListingPct-BuyerPct-ClosingPct),2)",
        # Period k = completion - 1 is the completion-th row of the table (k starts at 0).
        "BalanceAtCompletion": f"=INDEX({col('balanceEnd')},Completion)",
        "MonthlyBurn": (
            "=ROUND(BalanceAtCompletion*LoanRate/12,2)+ROUND(LandLoan*LandRate/12,2)"
            "+ROUND(TaxRate/12*(LandPrice+BuildSpent),2)+ROUND(InsuranceRate/12*EstimateTotal,2)"
            "+UtilAfter+StagingMonthly+UpkeepMonthly+HoaMonthly"
        ),
        "MaxDomNoExt": '=IF(Commitment=0,"",MAX(0,INT((Maturity-ListMonth)*MonthDays-ClosingDays)))',
        "CheckRows": "=PeriodCount-RowsInSchedule",
        "CheckSpend": f"=SUM({col('spend')})-BuildSpent",
        "CheckFunding": f"=SUM({col('equitySpend')})+SUM({col('loanDraw')})-BuildSpent",
        "CheckFlows": f"=EquityReturn-SUM({col('equityOut')})-Profit",
    }

    ws["A1"] = f"Golden deal - {name} scenario"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Every value below is a formula implementing the Reference cash-flow algorithm (dollars)."
    addresses: dict[str, str] = {}
    row = 3
    for key, label, fmt in SUMMARY:
        if key is None:  # section heading, after a blank row
            row += 1
            ws.cell(row=row, column=1, value=label).font = BOLD
            row += 1
            continue
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=f[key])
        if fmt:
            cell.number_format = fmt
        addresses[key] = f"B{row}"
        local_name(key, f"$B${row}")
        row += 1
    if row >= TABLE_HEADER_ROW:
        raise SystemExit(f"summary block overflows into the period table (row {row})")

    # Period table.
    for i, (key, header, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=TABLE_HEADER_ROW, column=i, value=header)
        cell.font = BOLD
    ws.cell(row=TABLE_HEADER_ROW - 1, column=1, value="Monthly schedule (one row per period k)").font = BOLD

    for r in range(first, last + 1):
        c = {key: f"{letter}{r}" for key, letter in COL.items()}
        prev = {key: f"{letter}{r - 1}" for key, letter in COL.items()}
        top = r == first
        k = c["k"]
        earlier_draws = "0" if top else f"SUM({COL['draw']}${first}:{prev['draw']})"
        row_formulas = {
            "k": "=0" if top else f"={prev['k']}+1",
            "ovBuildClose": f"=MAX(0,MIN({k}+1,SaleClose)-MAX({k},BuildStart))",
            "ovAll": f"=MAX(0,MIN({k}+1,SaleClose)-MAX({k},0))",
            "ovBuild": f"=MAX(0,MIN({k}+1,Completion)-MAX({k},BuildStart))",
            "ovAfter": f"=MAX(0,MIN({k}+1,SaleClose)-MAX({k},Completion))",
            "ovListed": f"=MAX(0,MIN({k}+1,Contract)-MAX({k},ListMonth))",
            "sine": (
                f"=IF(AND({k}>=BuildStart,{k}<BuildStart+BuildMonths),"
                f"SIN(PI()*({k}-BuildStart+0.5)/BuildMonths),0)"
            ),
            "draw": (
                f"=IF(AND({k}>=BuildStart,{k}<BuildStart+BuildMonths-1),"
                f"ROUND(HardActual*({c['sine']}/DrawWeightSum),2),"
                f"IF({k}=BuildStart+BuildMonths-1,HardActual-{earlier_draws},0))"
            ),
            "spend": f"=IF({k}=0,SoftActual,0)+{c['draw']}",
            "spendBefore": "=0" if top else f"={prev['spendThrough']}",
            "drawnBefore": "=0" if top else f"={prev['drawnThrough']}",
            "eqFirst": f"=MIN({c['spend']},MAX(0,EquityReq-{c['spendBefore']}))",
            "loanDraw": f"=MIN({c['spend']}-{c['eqFirst']},Commitment-{c['drawnBefore']})",
            "equitySpend": f"={c['eqFirst']}+({c['spend']}-{c['eqFirst']}-{c['loanDraw']})",
            "spendThrough": f"={c['spendBefore']}+{c['spend']}",
            "drawnThrough": f"={c['drawnBefore']}+{c['loanDraw']}",
            "balanceStart": "=0" if top else f"={prev['balanceEnd']}",
            "points": f"=IF(AND({k}=BuildStart,Commitment>0),ConstrPoints,0)",
            "drawFee": f"=IF({c['loanDraw']}>0,DrawFee,0)",
            "interest": f"=ROUND({c['balanceStart']}*LoanRate/12*{c['ovBuildClose']},2)",
            "interestPaid": f"=IF(CapFlag=1,0,{c['interest']})",
            "interestCap": f"=IF(CapFlag=1,{c['interest']},0)",
            "balanceEnd": f"={c['balanceStart']}+{c['loanDraw']}+{c['interestCap']}",
            "landInterest": f"=ROUND(LandLoan*LandRate/12*{c['ovAll']},2)",
            "extensionFee": f"=COUNTIF({col('extPeriod')},{k})*ExtFee",
            "tax": f"=ROUND(TaxRate/12*(LandPrice+{c['spendThrough']})*{c['ovAll']},2)",
            "insurance": f"=ROUND(InsuranceRate/12*EstimateTotal*{c['ovBuildClose']},2)",
            "utilities": f"=ROUND(UtilBuild*{c['ovBuild']}+UtilAfter*{c['ovAfter']},2)",
            "staging": f"=ROUND(StagingMonthly*{c['ovListed']},2)",
            "upkeep": f"=ROUND(UpkeepMonthly*{c['ovAfter']},2)",
            "hoa": f"=ROUND(HoaMonthly*{c['ovAll']},2)",
            "landOut": f"=IF({k}=0,LandPrice-LandLoan+LandClosing+LandPoints,0)",
            "equityOut": "=" + "+".join(
                c[key]
                for key in (
                    "landOut", "equitySpend", "points", "drawFee", "interestPaid", "landInterest",
                    "extensionFee", "tax", "insurance", "utilities", "staging", "upkeep", "hoa",
                )
            ),
            "cumEquity": f"={c['equityOut']}" if top else f"={prev['cumEquity']}+{c['equityOut']}",
            # Extension fee j (0..n-1) is paid at period floor(maturity + j x extensionMonths).
            # n <= N always, so the N table rows double as j = 0..N-1.
            "extJ": f"={k}",
            "extPeriod": f'=IF({c["extJ"]}<ExtCount,INT(Maturity+{c["extJ"]}*ExtensionMonths),"")',
        }
        for i, (key, _, fmt) in enumerate(COLUMNS, start=1):
            if key is None:
                continue
            cell = ws.cell(row=r, column=i, value=row_formulas[key])
            if fmt:
                cell.number_format = fmt

    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 18
    for key, letter in COL.items():
        if letter not in ("A", "B"):
            ws.column_dimensions[letter].width = 16
    ws.freeze_panes = ws.cell(row=first, column=2)
    return addresses


# ---------------------------------------------------------------------------------------------
# Recalculate and read back
# ---------------------------------------------------------------------------------------------


def recalculate(xlsx: Path, workdir: Path) -> Path:
    """Recalculates every formula with headless LibreOffice; returns the converted copy."""
    soffice = shutil.which("soffice")
    if soffice is None:
        raise SystemExit("LibreOffice (soffice) is required to recalculate the golden spreadsheet")
    profile = workdir / "lo-profile"
    outdir = workdir / "out"
    outdir.mkdir()
    cmd = [
        soffice,
        f"-env:UserInstallation={profile.as_uri()}",
        "--headless",
        "--calc",
        "--convert-to",
        "xlsx",
        "--outdir",
        str(outdir),
        str(xlsx),
    ]
    done = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    converted = outdir / xlsx.name
    if done.returncode != 0 or not converted.exists():
        raise SystemExit(f"LibreOffice conversion failed:\n{done.stdout}\n{done.stderr}")
    return converted


PERIOD_FIELDS = [
    ("spendCents", "spend"),
    ("equitySpendCents", "equitySpend"),
    ("loanDrawCents", "loanDraw"),
    ("drawFeeCents", "drawFee"),
    ("pointsCents", "points"),
    ("interestCents", "interest"),
    ("landLoanInterestCents", "landInterest"),
    ("extensionFeeCents", "extensionFee"),
    ("taxCents", "tax"),
    ("insuranceCents", "insurance"),
    ("utilitiesCents", "utilities"),
    ("stagingCents", "staging"),
    ("upkeepCents", "upkeep"),
    ("hoaCents", "hoa"),
    ("equityOutCents", "equityOut"),
]
SALE_FIELDS = [
    ("commissionsCents", "Commissions"),
    ("closingCents", "SaleClosing"),
    ("netCents", "NetProceeds"),
    ("loanPayoffCents", "Payoff"),
    ("equityReturnCents", "EquityReturn"),
]
METRIC_FIELDS = [
    ("totalProjectCostCents", "TotalProjectCost"),
    ("buildSpentCents", "BuildSpent"),
    ("financingCents", "Financing"),
    ("carryCents", "Carry"),
    ("sellingCostsCents", "SellingCosts"),
    ("profitCents", "Profit"),
    ("peakCashCents", "PeakCash"),
    ("breakEvenSalePriceCents", "BreakEvenPrice"),
    ("monthlyBurnWhileListedCents", "MonthlyBurn"),
]
TIME_FIELDS = [
    ("buildStart", "BuildStart"),
    ("completion", "Completion"),
    ("list", "ListMonth"),
    ("contract", "Contract"),
    ("saleClose", "SaleClose"),
    ("maturity", "Maturity"),
]


def read_scenario(ws: Worksheet, addresses: dict[str, str], n: int) -> dict:
    def summary(key: str) -> object:
        return ws[addresses[key]].value

    def number(key: str) -> float:
        v = summary(key)
        if not isinstance(v, (int, float)) or isinstance(v, bool):
            raise SystemExit(f"{ws.title}!{addresses[key]} ({key}): expected a number, got {v!r}")
        return v

    def whole(key: str) -> int:
        v = number(key)
        if v != int(v):
            raise SystemExit(f"{ws.title}!{addresses[key]} ({key}): expected a whole number, got {v!r}")
        return int(v)

    def money(key: str) -> int:
        return to_cents(summary(key), f"{ws.title}!{addresses[key]} ({key})")

    if whole("PeriodCount") != n:
        raise SystemExit(f"{ws.title}: sheet says N = {summary('PeriodCount')}, script wrote {n} rows")
    for check in ("CheckRows", "CheckSpend", "CheckFunding", "CheckFlows"):
        if abs(number(check)) > 1e-6:
            raise SystemExit(f"{ws.title}: check {check} is {summary(check)!r}, expected 0")

    periods = []
    period_ts = []
    for r in range(TABLE_FIRST_ROW, TABLE_FIRST_ROW + n):
        period_ts.append(ws[f"{COL['k']}{r}"].value)
        periods.append(
            {
                field: to_cents(ws[f"{COL[key]}{r}"].value, f"{ws.title}!{COL[key]}{r} ({key})")
                for field, key in PERIOD_FIELDS
            }
        )
    if period_ts != list(range(n)):
        raise SystemExit(f"{ws.title}: period indexes are {period_ts}, expected 0..{n - 1}")

    max_dom = summary("MaxDomNoExt")
    sale = {field: money(key) for field, key in SALE_FIELDS}
    metrics: dict[str, int | None] = {field: money(key) for field, key in METRIC_FIELDS}
    metrics["maxDomWithoutExtension"] = None if max_dom in (None, "") else whole("MaxDomNoExt")
    times = {field: number(key) for field, key in TIME_FIELDS}
    equity_flows = [{"t": k, "cents": -p["equityOutCents"]} for k, p in enumerate(periods)]
    equity_flows.append({"t": times["saleClose"], "cents": sale["equityReturnCents"]})

    # Integrity checks in exact integer cents: the sheet's period rows must reconcile to its totals.
    def total(field: str) -> int:
        return sum(p[field] for p in periods)

    reconciles = {
        "sum of spend = build spent": total("spendCents") == metrics["buildSpentCents"],
        "equity spend + loan draws = build spent": (
            total("equitySpendCents") + total("loanDrawCents") == metrics["buildSpentCents"]
        ),
        "sum of equity flows = profit": sum(f["cents"] for f in equity_flows) == metrics["profitCents"],
        "max cumulative equity out = peak cash": (
            max(accumulate(p["equityOutCents"] for p in periods)) == metrics["peakCashCents"]
        ),
    }
    for what, ok in reconciles.items():
        if not ok:
            raise SystemExit(f"{ws.title}: {what} does not hold")

    return {
        "inputs": {
            "priceCents": money("SalePrice"),
            "domDays": number("DomDays"),
            "overrunPct": number("Overrun"),
        },
        "times": times,
        "extensions": whole("ExtCount"),
        "periods": periods,
        "sale": sale,
        "metrics": metrics,
        "equityFlows": equity_flows,
    }


def main() -> None:
    deal = json.loads(INPUTS_JSON.read_text())
    base = {
        "priceCents": deal["overrides"]["basePriceCents"],
        "domDays": deal["overrides"]["baseDomDays"],
        "overrunPct": deal["scenarioDefs"]["base"]["overrunPct"],
    }
    scenarios = [("base", base), ("stress", STRESS)]

    wb = Workbook()
    scenario_rows = write_inputs_sheet(wb, deal, scenarios)
    layout: dict[str, tuple[int, dict[str, str]]] = {}
    for name, s in scenarios:
        n = period_count(deal["timeline"], s["domDays"])
        layout[name] = (n, write_schedule_sheet(wb, name, n, scenario_rows[name]))
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)

    with tempfile.TemporaryDirectory(prefix="golden-") as tmp:
        converted = recalculate(XLSX, Path(tmp))
        values = load_workbook(converted, data_only=True)
        expected = {
            "scenarios": {name: read_scenario(values[name], addresses, n) for name, (n, addresses) in layout.items()}
        }

    EXPECTED_JSON.write_text(json.dumps(expected, indent=2) + "\n")
    print(f"wrote {XLSX.relative_to(ROOT)}")
    print(f"wrote {EXPECTED_JSON.relative_to(ROOT)}")
    for name in layout:
        profit = expected["scenarios"][name]["metrics"]["profitCents"]
        sign = "-" if profit < 0 else ""
        print(f"{name} profit {sign}${abs(profit) / 100:,.2f}")


if __name__ == "__main__":
    main()
